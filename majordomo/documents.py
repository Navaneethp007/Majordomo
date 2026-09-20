"""Text out of documents that are not text files.

``read_file`` refuses binary outright, which is the honest floor: a PDF read as
UTF-8 returns replacement characters, the model summarises the noise, and you
get a confident answer about a document nobody read. This module is the step
above that floor — for the two formats worth the dependency.

Deliberately shaped like ``tts.py``: an adapter per format, a dict at the
bottom, and the optional imports done *inside* each adapter so a text-only
install never pays for them. Adding EPUB or RTF later is one function and one
entry.

── ON NAMING THE FAILURE ────────────────────────────────────────────────────
Three things go wrong here and they need different words, because the remedies
are different and the model acts on what it is told:

- **the extra is not installed** — name it, and it is one pip command away
- **the file is encrypted** — nothing in this process can help
- **the PDF is a scan with no text layer** — needs OCR, which is not offered

The third is the one that must never look like success. A scanned page yields
zero characters, and an empty string returned as "the document" reads exactly
like an empty document.
─────────────────────────────────────────────────────────────────────────────
"""
from __future__ import annotations

from pathlib import Path


class ExtractionError(Exception):
    """A document could not be turned into text. Always says which reason."""


class ExtractorMissing(ExtractionError):
    """The optional dependency for this format is not installed."""


#: Below this, a "successful" extraction is almost certainly a scanned page: the
#: text layer is absent and what came back is stray header junk. Chosen low
#: enough that a genuinely near-empty document still reads as one.
MIN_MEANINGFUL_CHARS = 16


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractorMissing(
            "reading PDFs needs an extra package — run: pip install majordomo[documents]"
        ) from exc

    from pypdf.errors import DependencyError, PdfReadError

    try:
        reader = PdfReader(str(path))

        if reader.is_encrypted:
            # `decrypt` **returns** a PasswordType and does not raise on a wrong
            # password — pypdf's own source carries a TODO about that. Wrapping
            # it in `except Exception` therefore caught nothing, execution fell
            # through to `reader.pages`, and the carefully worded message below
            # was unreachable: an encrypted PDF was reported as "could not be
            # read as a PDF" instead.
            #
            # The empty password opens the common "restricted but not secret"
            # case. A real password is not something to prompt for here.
            try:
                opened = reader.decrypt("")
            except DependencyError as exc:
                # AES without `cryptography` installed. Not a PdfReadError, so
                # it escaped this module entirely and surfaced as a generic
                # read failure with nothing actionable in it.
                raise ExtractionError(
                    f"{path.name} uses AES encryption, which needs another "
                    f"package — run: pip install cryptography"
                ) from exc

            if not opened:
                raise ExtractionError(
                    f"{path.name} is password-protected and cannot be opened here"
                )

        pages = [page.extract_text() or "" for page in reader.pages]

    except ExtractionError:
        raise
    except DependencyError as exc:
        raise ExtractionError(
            f"{path.name} needs another package to decode — run: "
            f"pip install cryptography"
        ) from exc
    except (PdfReadError, OSError, ValueError) as exc:
        raise ExtractionError(f"{path.name} could not be read as a PDF: {exc}") from exc

    text = "\n\n".join(p.strip() for p in pages if p.strip()).strip()
    if len(text) < MIN_MEANINGFUL_CHARS:
        raise ExtractionError(
            f"{path.name} has {len(reader.pages)} page(s) but almost no text — "
            f"it is most likely a scan. Reading it would need OCR, which is not "
            f"available here."
        )
    return text


def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractorMissing(
            "reading Word files needs an extra package — run: "
            "pip install majordomo[documents]"
        ) from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(
            f"{path.name} could not be read as a Word file: {exc}"
        ) from exc

    parts = [p.text for p in document.paragraphs]
    # Tables hold a lot of what people actually want out of a .docx, and
    # `paragraphs` does not include them.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    text = "\n".join(p for p in parts if p.strip()).strip()
    if not text:
        raise ExtractionError(f"{path.name} contains no text")
    return text


#: Suffix to adapter. The dict is the whole registry — see the module docstring.
EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
}


def can_extract(path: Path) -> bool:
    """Is this a format we know how to turn into text?

    Says nothing about whether the dependency is installed — that is answered by
    trying, so the error can name the extra rather than pretending the format is
    unknown.
    """
    return path.suffix.lower() in EXTRACTORS


def extract(path: Path) -> str:
    """The document's text.

    Raises:
        ExtractorMissing: the optional package is not installed.
        ExtractionError: encrypted, corrupt, or a scan with no text layer.
    """
    extractor = EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise ExtractionError(f"no way to extract text from {path.suffix or 'this file'}")
    return extractor(path)
